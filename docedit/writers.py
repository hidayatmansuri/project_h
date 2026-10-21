"""
Document exporters.

Writes edited document content preserving all layout, formatting, tables, lines,
and typography to the user's choice of output format:
- PDF: generated with embedded matching font, tables, borders, and lines via PyMuPDF Story.
- Word DOCX: generated with python-docx preserving native tables, header shading, borders, lines, alignments, and typography.
- Images (PNG, JPG, TIFF, BMP, WEBP): rendered via PyMuPDF with matching font, tables, and rules.
- Markdown: exported with markdown tables, divider rules, and headers.
- HTML: exported with clean embedded responsive CSS for tables, lines, and typography.
- Plain text: formatted with ASCII tables, divider lines, and text structure.
"""
import html
import io
import logging
import re
import unicodedata
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
from PIL import Image
import pymupdf

from .conf import EXTENSION_KINDS, conf
from .embedded_fonts import can_embedded_font_render_text, extract_all_embedded_fonts
from .fontmatch import is_non_latin_script
from .security import sanitize_document_html, sanitize_font_name

logger = logging.getLogger(__name__)


def normalize_text_for_diff(text: Any) -> str:
    """Normalizes text with NFC, replaces NBSP with space, and collapses whitespace."""
    if not text:
        return ""
    t = unicodedata.normalize("NFC", str(text))
    t = t.replace("\u00a0", " ")  # Non-breaking space
    return " ".join(t.split())


def is_block_truly_modified(b: Dict[str, Any]) -> bool:
    """
    Returns True only if the block has actually changed from its original extraction state.
    Prevents false-positives from focus, clicking, or whitespace differences.
    Never marks read-only or non-Latin blocks as modified to prevent exporting broken unshaped text.
    """
    if not b:
        return False
    if b.get("is_readonly"):
        return False
    curr_text = str(b.get("text", ""))
    orig_text = str(b.get("orig_text", ""))
    if is_non_latin_script(curr_text)[0] or (orig_text and is_non_latin_script(orig_text)[0]):
        return False
    if b.get("is_moved") or b.get("is_deleted") or b.get("is_new"):
        return True
    if not b.get("is_modified", False):
        return False
    if b.get("orig_text") is not None:
        clean_curr = re.sub(r"\n\s*\+\d+%\s*$", "", curr_text).strip()
        clean_orig = re.sub(r"\n\s*\+\d+%\s*$", "", orig_text).strip()
        curr_chars = re.sub(r"\s+", "", clean_curr)
        orig_chars = re.sub(r"\s+", "", clean_orig)
        if curr_chars != orig_chars and normalize_text_for_diff(clean_curr) != normalize_text_for_diff(clean_orig):
            return True
        if b.get("font_name") and b.get("orig_font_name"):
            f1 = re.sub(r"[\s\-_]|bold|italic|regular|mt|ps", "", str(b["font_name"]).lower())
            f2 = re.sub(r"[\s\-_]|bold|italic|regular|mt|ps", "", str(b["orig_font_name"]).lower())
            if f1 and f2 and f1 != f2:
                return True
        if b.get("font_size_pt") and b.get("orig_font_size_pt") and abs(b["font_size_pt"] - b["orig_font_size_pt"]) > 0.1:
            return True
        if b.get("orig_is_bold") is not None and bool(b.get("is_bold")) != bool(b.get("orig_is_bold")):
            return True
        if b.get("orig_is_italic") is not None and bool(b.get("is_italic")) != bool(b.get("orig_is_italic")):
            return True
        if b.get("orig_is_underline") is not None and bool(b.get("is_underline")) != bool(b.get("orig_is_underline")):
            return True
        return False
    return bool(b.get("is_modified", False))


def compute_redact_rect_for_block(
    b: Dict[str, Any],
    page_width: float = 595.0,
    page_height: float = 842.0,
    all_blocks: Optional[List[Dict[str, Any]]] = None,
    font_size_pt: float = 12.0,
) -> Optional[pymupdf.Rect]:
    """
    Computes a tight, neighbor-aware redaction rectangle for a modified or deleted text block.
    Returns None for newly added blocks (is_new=True) or read-only/non-Latin blocks.
    Redaction is strictly restricted to erasing the original text in orig_bbox without
    encroaching on adjacent neighbor text.
    """
    if b.get("is_new") or b.get("is_readonly"):
        return None
    curr_text = str(b.get("text", ""))
    orig_text = str(b.get("orig_text", ""))
    if is_non_latin_script(curr_text)[0] or (orig_text and is_non_latin_script(orig_text)[0]):
        return None
    bbox = b.get("bbox")
    orig_bbox = b.get("orig_bbox") or bbox
    if not orig_bbox or len(orig_bbox) != 4:
        return None

    rotation = int(b.get("rotation", 0))

    other_blocks = [
        o for o in (all_blocks or [])
        if o.get("id") != b.get("id") and o.get("orig_bbox") and len(o.get("orig_bbox")) == 4
    ]

    # Base redaction rectangle with 0.5pt margin to cover anti-aliasing edges
    redact_x0 = orig_bbox[0] - 0.5
    redact_y0 = orig_bbox[1] - 0.5
    redact_x1 = orig_bbox[2] + 0.5
    redact_y1 = orig_bbox[3] + 0.5

    # Clip against neighbor blocks so redaction NEVER wipes out adjacent labels or text
    for o in other_blocks:
        ob = o["orig_bbox"]
        # Check vertical overlap
        if not (ob[3] <= redact_y0 or ob[1] >= redact_y1):
            if ob[0] >= orig_bbox[0] and ob[0] < redact_x1:
                redact_x1 = min(redact_x1, ob[0] - 0.2)
            if ob[2] <= orig_bbox[2] and ob[2] > redact_x0:
                redact_x0 = max(redact_x0, ob[2] + 0.2)

        # Check horizontal overlap
        if not (ob[2] <= redact_x0 or ob[0] >= redact_x1):
            if ob[1] >= orig_bbox[1] and ob[1] < redact_y1:
                redact_y1 = min(redact_y1, ob[1] - 0.2)
            if ob[3] <= orig_bbox[3] and ob[3] > redact_y0:
                redact_y0 = max(redact_y0, ob[3] + 0.2)

    if redact_x0 >= redact_x1 or redact_y0 >= redact_y1:
        # Fall back to exact orig_bbox
        redact_x0 = orig_bbox[0]
        redact_y0 = orig_bbox[1]
        redact_x1 = orig_bbox[2]
        redact_y1 = orig_bbox[3]

    return pymupdf.Rect(redact_x0, redact_y0, redact_x1, redact_y1)


def get_median_background_color_ring(img: Image.Image, bbox: Any, ring_thickness: int = 4) -> Tuple[int, int, int]:
    """
    Computes the median RGB background color from a ring of pixels around the bounding box.
    Avoids sampling foreground text pixels inside the bounding box.
    """
    w, h = img.size
    ix0, iy0, ix1, iy1 = bbox
    ox0 = max(0, ix0 - ring_thickness)
    oy0 = max(0, iy0 - ring_thickness)
    ox1 = min(w, ix1 + ring_thickness)
    oy1 = min(h, iy1 + ring_thickness)

    arr = np.array(img)
    outer_crop = arr[oy0:oy1, ox0:ox1]

    # Create mask for the inner box
    mask = np.ones((oy1 - oy0, ox1 - ox0), dtype=bool)
    inner_y0 = max(0, iy0 - oy0)
    inner_y1 = min(oy1 - oy0, iy1 - oy0)
    inner_x0 = max(0, ix0 - ox0)
    inner_x1 = min(ox1 - ox0, ix1 - ox0)
    mask[inner_y0:inner_y1, inner_x0:inner_x1] = False

    ring_pixels = outer_crop[mask]
    if len(ring_pixels) == 0:
        return (255, 255, 255)

    median_rgb = np.median(ring_pixels, axis=0).astype(int)
    return tuple(int(c) for c in median_rgb[:3])



class DocumentHtmlBlockParser(HTMLParser):
    """
    Parses editor HTML into structured block elements preserving exact order:
    - headings (h1, h2, h3)
    - paragraphs (p)
    - divider lines / rules (hr)
    - tables (table with rows, cells, and headers)
    - list items (li)
    """
    def __init__(self):
        super().__init__()
        self.blocks = []
        self.in_table = False
        self.table_rows = []
        self.current_row = []
        self.current_cell_tag = None
        self.current_block_tag = None
        self.current_block_attrs = {}
        self.text_buf = []

    def handle_starttag(self, tag, attrs):
        t = tag.lower()
        attrs_dict = dict(attrs)

        if t == "table":
            self._flush_block()
            self.in_table = True
            self.table_rows = []
        elif self.in_table:
            if t == "tr":
                self.current_row = []
            elif t in ("td", "th"):
                self.current_cell_tag = t
                self.text_buf = []
            elif t == "a":
                href = attrs_dict.get("href", "")
                style_val = attrs_dict.get("style", "")
                self.text_buf.append(f'<a href="{html.escape(href)}" style="{html.escape(style_val)}">')
            elif t == "img":
                src = attrs_dict.get("src", "")
                alt = attrs_dict.get("alt", "")
                style_val = attrs_dict.get("style", "")
                self.text_buf.append(f'<img src="{src}" alt="{html.escape(alt)}" style="{html.escape(style_val)}"/>')
            elif t == "span":
                style_val = attrs_dict.get("style", "")
                if style_val:
                    self.text_buf.append(f'<span style="{style_val}">')
                else:
                    self.text_buf.append('<span>')
            elif t in ("b", "strong"):
                self.text_buf.append("<strong>")
            elif t in ("i", "em"):
                self.text_buf.append("<em>")
            elif t == "u":
                self.text_buf.append("<u>")
            elif t == "br":
                self.text_buf.append("<br/>")
        elif t in ("h1", "h2", "h3", "p", "li", "blockquote"):
            self._flush_block()
            self.current_block_tag = t
            self.current_block_attrs = attrs_dict
            self.text_buf = []
        elif t == "hr":
            self._flush_block()
            self.blocks.append({"type": "hr"})
        elif t == "a":
            href = attrs_dict.get("href", "")
            style_val = attrs_dict.get("style", "")
            self.text_buf.append(f'<a href="{html.escape(href)}" style="{html.escape(style_val)}">')
        elif t == "img":
            src = attrs_dict.get("src", "")
            alt = attrs_dict.get("alt", "")
            style_val = attrs_dict.get("style", "")
            self.text_buf.append(f'<img src="{src}" alt="{html.escape(alt)}" style="{html.escape(style_val)}"/>')
        elif t == "span":
            style_val = attrs_dict.get("style", "")
            if style_val:
                self.text_buf.append(f'<span style="{style_val}">')
            else:
                self.text_buf.append('<span>')
        elif t in ("b", "strong"):
            self.text_buf.append("<strong>")
        elif t in ("i", "em"):
            self.text_buf.append("<em>")
        elif t == "u":
            self.text_buf.append("<u>")
        elif t == "br":
            self.text_buf.append("<br/>")

    def handle_endtag(self, tag):
        t = tag.lower()
        if t == "table":
            self.in_table = False
            if self.table_rows:
                self.blocks.append({"type": "table", "rows": self.table_rows})
            self.table_rows = []
        elif self.in_table:
            if t in ("td", "th"):
                cell_html = "".join(self.text_buf).strip()
                is_hdr = (self.current_cell_tag == "th")
                self.current_row.append({"is_header": is_hdr, "html": cell_html})
                self.text_buf = []
                self.current_cell_tag = None
            elif t == "tr":
                if self.current_row:
                    self.table_rows.append(self.current_row)
                    self.current_row = []
            elif t == "a":
                self.text_buf.append("</a>")
            elif t == "span":
                self.text_buf.append("</span>")
            elif t in ("b", "strong"):
                self.text_buf.append("</strong>")
            elif t in ("i", "em"):
                self.text_buf.append("</em>")
            elif t == "u":
                self.text_buf.append("</u>")
        elif t in ("h1", "h2", "h3", "p", "li", "blockquote"):
            self._flush_block()
        elif t == "a":
            self.text_buf.append("</a>")
        elif t == "span":
            self.text_buf.append("</span>")
        elif t in ("b", "strong"):
            self.text_buf.append("</strong>")
        elif t in ("i", "em"):
            self.text_buf.append("</em>")
        elif t == "u":
            self.text_buf.append("</u>")

    def handle_data(self, data):
        self.text_buf.append(data)

    def _flush_block(self):
        if self.current_block_tag:
            content = "".join(self.text_buf).strip()
            if content:
                style_str = self.current_block_attrs.get("style", "").lower()
                align = None
                if "text-align: center" in style_str:
                    align = "center"
                elif "text-align: right" in style_str:
                    align = "right"
                elif "text-align: justify" in style_str:
                    align = "justify"
                elif "text-align: left" in style_str:
                    align = "left"

                font_match = re.search(r'font-family:\s*[\'"]?([^\'";]+)[\'"]?', style_str, re.IGNORECASE)
                b_font_name = font_match.group(1).split(",")[0].strip().strip("'\"") if font_match else None
                size_match = re.search(r'font-size:\s*([\d\.]+)pt', style_str, re.IGNORECASE)
                b_font_size = float(size_match.group(1)) if size_match else None

                self.blocks.append({
                    "type": self.current_block_tag,
                    "html": content,
                    "align": align,
                    "font_name": b_font_name,
                    "font_size_pt": b_font_size,
                })
            self.current_block_tag = None
            self.current_block_attrs = {}
            self.text_buf = []

    def close(self):
        super().close()
        self._flush_block()


def parse_html_blocks(html_content: str) -> List[Dict[str, Any]]:
    """Strips editor artifacts and parses HTML into structured block elements."""
    # Clean fuzzy highlights
    clean_html = re.sub(r'<span[^>]*class="[^"]*docedit-fuzzy[^"]*"[^>]*>(.*?)</span>', r'\1', html_content)
    clean_html = re.sub(r'data-[a-zA-Z0-9_\-]+="[^"]*"', '', clean_html)
    parser = DocumentHtmlBlockParser()
    parser.feed(clean_html)
    parser.close()
    return parser.blocks


def _write_pdf_layout_preserved(
    pages_data: List[Dict[str, Any]],
    original_file_path: Optional[Path],
    output_path: Path,
    font_name: str,
    font_path: Optional[str],
    font_size_pt: float,
    work_dir: Optional[Path] = None,
) -> Path:
    """
    Exports PDF while preserving 100% of the original document layout, tables,
    vector lines, borders, and graphics. Applies in-place redaction and text insertion
    only to modified text blocks.
    """
    import pymupdf

    from .font_catalog import is_font_available_locally, _find_system_font_file, find_font_variant_file, resolve_document_font

    def get_font_kwargs_for_block(
        b_name: str,
        preferred_path: Optional[str],
        is_bold: bool = False,
        is_italic: bool = False
    ) -> Dict[str, Any]:
        # 1. Search for matching font variant file
        var_p = find_font_variant_file(b_name, is_bold=is_bold, is_italic=is_italic, base_font_path=preferred_path)
        if var_p and Path(var_p).is_file():
            try:
                fobj = pymupdf.Font(fontfile=str(var_p))
                reg_name = fobj.name.replace(" ", "-")
            except Exception:
                reg_name = f"{b_name}-{'Bold' if is_bold else ''}{'Italic' if is_italic else ''}".strip("-")
            return {"fontname": reg_name, "fontfile": str(var_p)}

        # 2. Try resolving via font catalog (system lookup or free download)
        try:
            res = resolve_document_font(b_name)
            if res and res.get("font_path") and Path(res["font_path"]).is_file():
                var_p2 = find_font_variant_file(res.get("resolved_name", b_name), is_bold=is_bold, is_italic=is_italic, base_font_path=res["font_path"])
                fp = var_p2 if (var_p2 and Path(var_p2).is_file()) else res["font_path"]
                try:
                    fobj = pymupdf.Font(fontfile=str(fp))
                    reg_name = fobj.name.replace(" ", "-")
                except Exception:
                    reg_name = f"{b_name}-{'Bold' if is_bold else ''}{'Italic' if is_italic else ''}".strip("-")
                return {"fontname": reg_name, "fontfile": str(fp)}
        except Exception:
            pass

        # 3. PyMuPDF built-in Base 14 fonts with exact style variants
        low = (b_name or "").lower()
        if any(s in low for s in ("times", "serif", "garamond", "georgia", "baskerville", "caslon", "cambria", "book")):
            if is_bold and is_italic:
                return {"fontname": "tibi"}
            elif is_bold:
                return {"fontname": "tibo"}
            elif is_italic:
                return {"fontname": "tiit"}
            return {"fontname": "tiro"}
        elif any(s in low for s in ("mono", "courier", "console", "consolas")):
            if is_bold and is_italic:
                return {"fontname": "cobi"}
            elif is_bold:
                return {"fontname": "cobo"}
            elif is_italic:
                return {"fontname": "coit"}
            return {"fontname": "couri"}
        else:  # Sans serif (helv, arial, calibri, segoe, etc.)
            if is_bold and is_italic:
                return {"fontname": "hebi"}
            elif is_bold:
                return {"fontname": "hebo"}
            elif is_italic:
                return {"fontname": "heit"}
            return {"fontname": "helv"}

    embedded_catalog = None
    if original_file_path and original_file_path.is_file() and original_file_path.suffix.lower() == ".pdf":
        try:
            temp_doc = pymupdf.open(str(original_file_path))
            embedded_catalog = extract_all_embedded_fonts(temp_doc, session_dir=work_dir)
            temp_doc.close()
        except Exception as e:
            logger.debug("Failed extracting embedded fonts for writer: %s", e)

    # compute_redact_rect_for_block is defined at module-level

    def insert_modified_block_into_page(page, b, page_width: float, page_height: float, all_blocks: Optional[List[Dict[str, Any]]] = None):
        if b.get("is_readonly") or is_non_latin_script(str(b.get("text", "")))[0]:
            logger.warning("Skipping text insertion for non-Latin block %s to avoid broken text", b.get("id"))
            return
        bbox = b.get("bbox")
        orig_bbox = b.get("orig_bbox") or bbox
        if not orig_bbox or len(orig_bbox) != 4:
            return
        f_size = float(b.get("font_size_pt") or font_size_pt)
        b_fname = b.get("font_name") or font_name
        is_b = bool(b.get("is_bold", False))
        is_it = bool(b.get("is_italic", False))
        text_str = re.sub(r"\n\s*\+\d+%\s*$", "", str(b.get("text", ""))).strip()
        rotation = int(b.get("rotation", 0))
        orig_origin = b.get("origin")

        # Check if an embedded font exists that covers this text without restriction
        has_fontfile = False
        font_file_path = None
        base_fontname = "helv"
        b_fkw = None

        if embedded_catalog:
            clean_k = b_fname.lower()
            cand = embedded_catalog.get("by_clean_name", {}).get(clean_k)
            if not cand:
                for k, v in embedded_catalog.get("by_clean_name", {}).items():
                    if k in clean_k or clean_k in k:
                        cand = v
                        break
            if cand and cand.get("file_path") and Path(cand["file_path"]).is_file():
                if can_embedded_font_render_text(cand, text_str):
                    f_path = cand["file_path"]
                    clean_reg = re.sub(r"[^a-zA-Z0-9_-]", "", cand["clean_name"])
                    b_fkw = {"fontname": f"{clean_reg}-Emb", "fontfile": str(f_path)}
                    has_fontfile = True
                    font_file_path = str(f_path)
                    base_fontname = f"{clean_reg}-Emb"

        if not b_fkw and b.get("font_info", {}).get("font_path"):
            cand_fp = Path(b["font_info"]["font_path"])
            if cand_fp.is_file():
                safe_tag = re.sub(r"[^\w]", "", b_fname or "Font")
                b_fkw = {"fontname": f"{safe_tag}-Res", "fontfile": str(cand_fp)}
                has_fontfile = True
                font_file_path = str(cand_fp)
                base_fontname = f"{safe_tag}-Res"

        if not b_fkw:
            b_fkw = get_font_kwargs_for_block(b_fname, font_path, is_bold=is_b, is_italic=is_it)
            has_fontfile = "fontfile" in b_fkw
            font_file_path = b_fkw.get("fontfile")
            base_fontname = b_fkw.get("fontname", "helv")

        # Extract block text color
        col = b.get("color_rgb")
        if not col and b.get("color"):
            hex_s = str(b.get("color")).lstrip("#")
            if len(hex_s) == 6:
                try:
                    col = (int(hex_s[0:2], 16) / 255.0, int(hex_s[2:4], 16) / 255.0, int(hex_s[4:6], 16) / 255.0)
                except Exception:
                    col = (0.0, 0.0, 0.0)
        if not col:
            col = (0.0, 0.0, 0.0)

        safe_text = (
            text_str.replace("\u2013", "-")
            .replace("\u2014", "--")
            .replace("\u2018", "'")
            .replace("\u2019", "'")
            .replace("\u201c", '"')
            .replace("\u201d", '"')
        )

        # Compute safe built-in fallback fontname (Base-14)
        low_fname = (b_fname or "").lower()
        if any(s in low_fname for s in ("times", "serif", "garamond", "georgia", "baskerville", "caslon", "cambria", "book")):
            if is_b and is_it:
                base14_fallback = "tibi"
            elif is_b:
                base14_fallback = "tibo"
            elif is_it:
                base14_fallback = "tiit"
            else:
                base14_fallback = "tiro"
        elif any(s in low_fname for s in ("mono", "courier", "console", "consolas")):
            if is_b and is_it:
                base14_fallback = "cobi"
            elif is_b:
                base14_fallback = "cobo"
            elif is_it:
                base14_fallback = "coit"
            else:
                base14_fallback = "couri"
        else:
            if is_b and is_it:
                base14_fallback = "hebi"
            elif is_b:
                base14_fallback = "hebo"
            elif is_it:
                base14_fallback = "heit"
            else:
                base14_fallback = "helv"

        active_fn = base14_fallback
        font_registered = False
        if has_fontfile and font_file_path and Path(font_file_path).is_file():
            is_ascii = all(ord(c) < 128 for c in text_str)
            reg_fontname = f"{base_fontname}{'-Simp' if is_ascii else '-Uni'}"
            try:
                page.insert_font(fontname=reg_fontname, fontfile=str(font_file_path), set_simple=is_ascii)
                active_fn = reg_fontname
                font_registered = True
            except Exception as e:
                logger.debug("insert_font note for %s: %s; using %s", reg_fontname, e, base14_fallback)
                active_fn = base14_fallback
                font_registered = False
        else:
            active_fn = base14_fallback
            font_registered = False

        out_txt = text_str if font_registered else safe_text

        def safe_insert(pt, txt, font_sz, fn, clr, rotate=0, rot=None):
            if rot is not None:
                rotate = rot
            try:
                page.insert_text(pt, txt, fontsize=font_sz, fontname=fn, color=clr, rotate=rotate)
            except Exception as ex1:
                logger.warning("insert_text failed with font %s: %s; falling back to %s", fn, ex1, base14_fallback)
                try:
                    page.insert_text(pt, safe_text, fontsize=font_sz, fontname=base14_fallback, color=clr, rotate=rotate)
                except Exception as ex2:
                    logger.error("insert_text Base-14 fallback failed: %s; falling back to helv", ex2)
                    page.insert_text(pt, safe_text, fontsize=font_sz, fontname="helv", color=clr, rotate=rotate)

        # Measure text length
        tl = None
        if has_fontfile and font_file_path:
            try:
                fobj = pymupdf.Font(fontfile=str(font_file_path))
                first_line = out_txt.split("\n")[0] if "\n" in out_txt else out_txt
                tl = fobj.text_length(first_line, fontsize=f_size)
            except Exception:
                tl = None
        if tl is None:
            try:
                first_line = out_txt.split("\n")[0] if "\n" in out_txt else out_txt
                tl = pymupdf.get_text_length(first_line, fontname=active_fn, fontsize=f_size)
            except Exception:
                try:
                    tl = pymupdf.get_text_length(first_line, fontname=base14_fallback, fontsize=f_size)
                except Exception:
                    tl = (orig_bbox[2] - orig_bbox[0]) if orig_bbox and len(orig_bbox) == 4 else 50.0

        is_multi_line = ("\n" in text_str) or bool(b.get("is_paragraph")) or bool(b.get("lines") and len(b["lines"]) > 1)

        if rotation == 90:
            x_pos = orig_origin[0] if (orig_origin and len(orig_origin) == 2) else (orig_bbox[2] - max(1.0, f_size * 0.15))
            y_pos = orig_origin[1] if (orig_origin and len(orig_origin) == 2) else orig_bbox[3]
            safe_insert(pymupdf.Point(x_pos, y_pos), out_txt, f_size, active_fn, col, rotate=90)
        elif rotation == 270:
            x_pos = orig_origin[0] if (orig_origin and len(orig_origin) == 2) else (orig_bbox[0] + max(1.0, f_size * 0.15))
            y_pos = orig_origin[1] if (orig_origin and len(orig_origin) == 2) else orig_bbox[1]
            safe_insert(pymupdf.Point(x_pos, y_pos), out_txt, f_size, active_fn, col, rotate=270)
        elif rotation == 180:
            x_pos = orig_origin[0] if (orig_origin and len(orig_origin) == 2) else orig_bbox[2]
            y_pos = orig_origin[1] if (orig_origin and len(orig_origin) == 2) else (orig_bbox[1] + max(1.5, f_size * 0.18))
            safe_insert(pymupdf.Point(x_pos, y_pos), out_txt, f_size, active_fn, col, rotate=180)
        else:
            # rotation == 0
            if is_multi_line:
                # Multi-line paragraphs: prefer insert_htmlbox with scale_low=0.7 so text reflows cleanly
                box_rect = pymupdf.Rect(b.get("bbox") or orig_bbox)
                css_str = ""
                if font_file_path and Path(font_file_path).is_file():
                    f_url = Path(font_file_path).as_posix()
                    css_str = f"@font-face {{ font-family: 'BlockFont'; src: url('{f_url}'); }}\n"
                    css_fn = "'BlockFont', sans-serif"
                else:
                    safe_bfname = html.escape(sanitize_font_name(b_fname), quote=True)
                    is_serif = any(s in safe_bfname.lower() for s in ["times", "serif", "garamond", "georgia", "baskerville", "caslon", "cambria", "book"])
                    css_fn = f"'{safe_bfname}', {'serif' if is_serif else 'sans-serif'}"

                align = b.get("align", "left")
                hex_col = f"#{int(col[0]*255):02x}{int(col[1]*255):02x}{int(col[2]*255):02x}"
                escaped_body = html.escape(text_str).replace("\n", "<br/>")
                html_box_str = f'<div style="font-family: {css_fn}; font-size: {f_size}pt; color: {hex_col}; text-align: {align}; margin:0; padding:0; line-height: 1.2;">{escaped_body}</div>'
                try:
                    spare_h, scale_factor = page.insert_htmlbox(box_rect, html_box_str, css=css_str, scale_low=0.7)
                    if spare_h < 0:
                        logger.warning("Block %s text overflowed bounding box by %.1f pt (scale factor %.2f)", b.get("id"), -spare_h, scale_factor)
                        b["overflow_warn"] = True
                        b["overflow_pct"] = round((-spare_h / max(1.0, box_rect.height)) * 100)
                        test_drawn = page.get_text("text", clip=box_rect).strip()
                        if not test_drawn:
                            # insert_htmlbox failed to fit anything, fall back to safe_insert lines
                            lines_to_draw = text_str.split("\n")
                            lh = max(f_size * 1.15, box_rect.height / max(1, len(lines_to_draw)))
                            cur_y = box_rect.y0 + f_size
                            for ltxt in lines_to_draw:
                                safe_insert(pymupdf.Point(box_rect.x0, cur_y), ltxt, f_size, active_fn, col, rotate=0)
                                cur_y += lh
                except Exception as ex_box:
                    logger.warning("insert_htmlbox failed: %s; falling back to safe_insert", ex_box)
                    safe_insert(pymupdf.Point(box_rect.x0, box_rect.y0 + f_size), out_txt, f_size, active_fn, col, rotate=0)
            else:
                # Single-line: insert at stored baseline origin!
                box_rect = pymupdf.Rect(b.get("bbox") or orig_bbox)
                if tl is not None and box_rect.width > 0 and tl > box_rect.width * 1.05:
                    logger.warning("Block %s text width (%.1f pt) overflowed bounding box width (%.1f pt)", b.get("id"), tl, box_rect.width)
                    b["overflow_warn"] = True
                    b["overflow_pct"] = round(((tl - box_rect.width) / box_rect.width) * 100)

                is_right_aligned = (
                    orig_bbox[2] > page_width * 0.70
                    and orig_bbox[0] > page_width * 0.35
                    and (orig_bbox[2] >= (bbox[2] if bbox and len(bbox) == 4 else orig_bbox[2]) - 6)
                )
                if is_right_aligned and tl is not None:
                    x_pos = orig_bbox[2] - tl
                elif orig_origin and len(orig_origin) == 2 and not b.get("is_moved", False):
                    x_pos = orig_origin[0]
                else:
                    x_pos = bbox[0] if bbox and len(bbox) == 4 else orig_bbox[0]

                if orig_origin and len(orig_origin) == 2 and not b.get("is_moved", False):
                    y_pos = orig_origin[1]
                else:
                    rect = pymupdf.Rect(bbox if bbox and len(bbox) == 4 else orig_bbox)
                    y_pos = rect.y1 - max(1.5, rect.height * 0.18)

                safe_insert(pymupdf.Point(x_pos, y_pos), out_txt, f_size, active_fn, col, rotate=0)

        if b.get("is_underline", False):
            try:
                rect = pymupdf.Rect(orig_bbox)
                ul_y = min(rect.y1, (orig_origin[1] if orig_origin else rect.y1) + max(1.0, f_size * 0.12))
                stroke_w = max(0.6, f_size * 0.055)
                ul_x = orig_origin[0] if orig_origin else orig_bbox[0]
                page.draw_line(pymupdf.Point(ul_x, ul_y), pymupdf.Point(ul_x + (tl or 30), ul_y), color=col, width=stroke_w)
            except Exception as e:
                logger.debug("Failed drawing underline: %s", e)

    # Try modifying original PDF directly if available
    if original_file_path and original_file_path.is_file() and original_file_path.suffix.lower() == ".pdf":
        try:
            doc = pymupdf.open(str(original_file_path))
            for p_info in pages_data:
                p_num = p_info.get("page_num", 1)
                if not (0 <= p_num - 1 < len(doc)):
                    continue
                page = doc[p_num - 1]
                p_w = page.rect.width
                p_h = page.rect.height
                page_blocks = p_info.get("blocks", [])

                # Only process truly modified blocks
                modified_blocks = [b for b in page_blocks if is_block_truly_modified(b)]
                if not modified_blocks:
                    continue  # Untouched page: pixel-identical, 0 changes

                # 1. Queue redactions with fill=False for all modified blocks on this page
                for b in modified_blocks:
                    r = compute_redact_rect_for_block(b, p_w, p_h, page_blocks)
                    if r:
                        page.add_redact_annot(r, fill=False)

                # 2. Apply redactions ONCE per page preserving line art, vector tables, and images
                page.apply_redactions(
                    images=pymupdf.PDF_REDACT_IMAGE_NONE,
                    graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                )

                # 3. Insert replacement text for each modified block
                for b in modified_blocks:
                    if not b.get("is_deleted", False):
                        insert_modified_block_into_page(page, b, p_w, p_h, all_blocks=page_blocks)

            doc.save(str(output_path), deflate=True)
            doc.close()
            return output_path
        except Exception as e:
            logger.warning("Failed direct PDF redaction, falling back to image compilation: %s", e)

    # Fallback: compile page background images with edited text overlay
    doc = pymupdf.open()
    for p_info in pages_data:
        p_num = p_info.get("page_num", 1)
        w = p_info.get("width", 595.0)
        h = p_info.get("height", 842.0)
        page = doc.new_page(width=w, height=h)

        # Place background image if available in work_dir
        if work_dir:
            img_p = work_dir / f"page_{p_num}.png"
            if img_p.is_file():
                page.insert_image(pymupdf.Rect(0, 0, w, h), filename=str(img_p))

        fallback_blocks = p_info.get("blocks", [])
        modified_blocks = [b for b in fallback_blocks if is_block_truly_modified(b)]
        if modified_blocks:
            for b in modified_blocks:
                r = compute_redact_rect_for_block(b, w, h, fallback_blocks)
                if r:
                    page.add_redact_annot(r, fill=False)

            page.apply_redactions(
                images=pymupdf.PDF_REDACT_IMAGE_NONE,
                graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
            )

            for b in modified_blocks:
                if not b.get("is_deleted", False):
                    insert_modified_block_into_page(page, b, w, h, all_blocks=fallback_blocks)

    doc.save(str(output_path), deflate=True)
    doc.close()
    return output_path


def _write_image_layout_preserved(
    pages_data: List[Dict[str, Any]],
    work_dir: Optional[Path],
    output_path: Path,
    image_ext: str,
    font_name: str,
    font_path: Optional[str],
    font_size_pt: float,
) -> Path:
    """
    Renders edited text in-place directly onto the page images, preserving original
    page borders, tables, graphics, and background texture.
    """
    from PIL import Image, ImageDraw, ImageFont
    from .font_catalog import find_font_variant_file, resolve_document_font

    pil_font = None
    if font_path and Path(font_path).is_file():
        try:
            pil_font = ImageFont.truetype(str(font_path), size=int(font_size_pt * 2.5))
        except Exception:
            pil_font = None

    rendered_pages = []
    for p_info in pages_data:
        p_num = p_info.get("page_num", 1)
        base_img = None

        if work_dir:
            p_img_path = work_dir / f"page_{p_num}.png"
            if p_img_path.is_file():
                base_img = Image.open(str(p_img_path)).convert("RGB")

        if base_img is None:
            w = int(p_info.get("width", 800))
            h = int(p_info.get("height", 1100))
            base_img = Image.new("RGB", (w, h), (255, 255, 255))

        draw = ImageDraw.Draw(base_img)
        w_img, h_img = base_img.size
        p_w = p_info.get("width", w_img)
        p_h = p_info.get("height", h_img)
        scale_x = w_img / max(1.0, float(p_w))
        scale_y = h_img / max(1.0, float(p_h))

        for b in p_info.get("blocks", []):
            if is_block_truly_modified(b):
                bbox = b.get("bbox")
                if bbox and len(bbox) == 4:
                    x0, y0, x1, y1 = bbox
                    ix0 = max(0, int(x0 * scale_x))
                    iy0 = max(0, int(y0 * scale_y))
                    ix1 = min(w_img, int(x1 * scale_x))
                    iy1 = min(h_img, int(y1 * scale_y))

                    # Compute ring background median color around the box (avoids interior text)
                    if not b.get("is_new"):
                        bg_color = get_median_background_color_ring(base_img, (ix0, iy0, ix1, iy1), ring_thickness=4)
                        draw.rectangle([ix0, iy0, ix1, iy1], fill=bg_color)

                    if b.get("is_deleted"):
                        continue

                    if b.get("size_src"):
                        b_size = int(round(b["size_src"] * scale_y))
                    elif p_info.get("unit") == "px":
                        p_dpi = p_info.get("dpi") or 150.0
                        b_size = int(round(b.get("font_size_pt", font_size_pt) * (p_dpi / 72.0) * scale_y))
                    else:
                        b_size = int(round(b.get("font_size_pt", font_size_pt) * scale_y))
                    b_size = max(6, b_size)

                    b_fname = b.get("font_name") or font_name
                    is_b = b.get("is_bold", False)
                    is_it = b.get("is_italic", False)
                    b_fpath = find_font_variant_file(b_fname, is_bold=is_b, is_italic=is_it, base_font_path=font_path)
                    if not b_fpath:
                        try:
                            res = resolve_document_font(b_fname)
                            if res and res.get("font_path") and Path(res["font_path"]).is_file():
                                b_fpath = find_font_variant_file(res.get("resolved_name", b_fname), is_bold=is_b, is_italic=is_it, base_font_path=res["font_path"]) or res["font_path"]
                        except Exception:
                            pass

                    cur_font = None
                    if b_fpath and Path(b_fpath).is_file():
                        try:
                            cur_font = ImageFont.truetype(str(b_fpath), size=max(10, b_size))
                        except Exception:
                            cur_font = pil_font
                    elif font_path and Path(font_path).is_file():
                        try:
                            cur_font = ImageFont.truetype(str(font_path), size=max(10, b_size))
                        except Exception:
                            cur_font = pil_font
                    else:
                        cur_font = pil_font

                    fill_col = (20, 20, 20)
                    if b.get("color_rgb"):
                        fill_col = tuple(int(round(c * 255)) for c in b["color_rgb"])
                    elif b.get("color"):
                        hex_c = str(b["color"]).lstrip("#")
                        if len(hex_c) == 6:
                            try:
                                fill_col = (int(hex_c[0:2], 16), int(hex_c[2:4], 16), int(hex_c[4:6], 16))
                            except Exception:
                                fill_col = (20, 20, 20)

                    # Baseline origin anchoring (anchor="ls" in PIL)
                    if b.get("origin") and len(b["origin"]) == 2:
                        ix_pos = int(round(b["origin"][0] * scale_x))
                        iy_pos = int(round(b["origin"][1] * scale_y))
                    else:
                        ix_pos = ix0
                        iy_pos = int(round(iy1 - (iy1 - iy0) * 0.15))

                    draw.text((ix_pos, iy_pos), b.get("text", ""), fill=fill_col, font=cur_font, anchor="ls")

        rendered_pages.append(base_img)

    fmt_map = {
        "jpg": "JPEG",
        "jpeg": "JPEG",
        "png": "PNG",
        "tif": "TIFF",
        "tiff": "TIFF",
        "bmp": "BMP",
        "webp": "WEBP",
    }
    save_fmt = fmt_map.get(image_ext.lower(), "PNG")

    if not rendered_pages:
        rendered_pages = [Image.new("RGB", (600, 800), (255, 255, 255))]

    if save_fmt == "TIFF" and len(rendered_pages) > 1:
        rendered_pages[0].save(
            str(output_path),
            format=save_fmt,
            save_all=True,
            append_images=rendered_pages[1:],
            dpi=(conf.OCR_DPI, conf.OCR_DPI)
        )
    else:
        rendered_pages[0].save(str(output_path), format=save_fmt, dpi=(conf.OCR_DPI, conf.OCR_DPI))

    return output_path


def _write_searchable_pdf(
    pages_data: List[Dict[str, Any]],
    work_dir: Optional[Path],
    original_file_path: Optional[Path],
    output_path: Path,
    font_name: str,
    font_path: Optional[str],
    font_size_pt: float,
) -> Path:
    """
    Exports scanned pages as a Searchable PDF:
    - Places the full-resolution page scan image as the visible page background.
    - If any text blocks were modified, erases and updates them cleanly on the background image.
    - Overlays recognized/edited text as an invisible selectable text layer (render_mode=3)
      positioned from line/word bboxes or baseline origin.
    """
    import pymupdf

    doc = pymupdf.open()
    for p_info in pages_data:
        p_num = p_info.get("page_num", 1)
        w = float(p_info.get("width", 595.0))
        h = float(p_info.get("height", 842.0))
        page = doc.new_page(width=w, height=h)

        page_blocks = p_info.get("blocks", [])
        has_modified = any(is_block_truly_modified(b) for b in page_blocks)

        # 1. Background Image
        bg_image_path = None
        if work_dir:
            p_img = work_dir / f"page_{p_num}.png"
            if p_img.is_file():
                bg_image_path = p_img

        if not bg_image_path and original_file_path and original_file_path.is_file():
            if original_file_path.suffix.lower() in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
                bg_image_path = original_file_path

        if has_modified and bg_image_path:
            # Render edited image with modified blocks erased via median ring and redrawn
            temp_img_out = (work_dir or Path(output_path.parent)) / f"temp_searchable_p{p_num}.png"
            _write_image_layout_preserved(
                pages_data=[p_info],
                work_dir=work_dir,
                output_path=temp_img_out,
                image_ext="png",
                font_name=font_name,
                font_path=font_path,
                font_size_pt=font_size_pt,
            )
            page.insert_image(pymupdf.Rect(0, 0, w, h), filename=str(temp_img_out))
            try:
                temp_img_out.unlink()
            except Exception:
                pass
        elif bg_image_path:
            page.insert_image(pymupdf.Rect(0, 0, w, h), filename=str(bg_image_path))

        # 2. Invisible Searchable Text Layer (render_mode=3)
        for b in page_blocks:
            if b.get("is_deleted", False):
                continue
            btxt = b.get("text", "")
            if not btxt.strip():
                continue

            lines = b.get("lines")
            if lines and len(lines) > 0:
                for line in lines:
                    ltxt = line.get("text", "")
                    if not ltxt.strip():
                        continue
                    lsz = float(line.get("font_size_pt") or b.get("font_size_pt") or font_size_pt)
                    lorigin = line.get("origin")
                    lbbox = line.get("bbox")
                    lrot = int(line.get("rotation") or b.get("rotation") or 0)
                    if lorigin and len(lorigin) == 2:
                        page.insert_text(pymupdf.Point(lorigin[0], lorigin[1]), ltxt, fontsize=lsz, render_mode=3, rotate=lrot)
                    elif lbbox and len(lbbox) == 4:
                        page.insert_textbox(pymupdf.Rect(lbbox), ltxt, fontsize=lsz, render_mode=3, rotate=lrot)
            else:
                bsz = float(b.get("font_size_pt") or font_size_pt)
                borigin = b.get("origin")
                bbbox = b.get("bbox")
                brot = int(b.get("rotation") or 0)
                if borigin and len(borigin) == 2 and "\n" not in btxt:
                    page.insert_text(pymupdf.Point(borigin[0], borigin[1]), btxt, fontsize=bsz, render_mode=3, rotate=brot)
                elif bbbox and len(bbbox) == 4:
                    page.insert_textbox(pymupdf.Rect(bbbox), btxt, fontsize=bsz, render_mode=3, rotate=brot)

    doc.save(str(output_path), deflate=True)
    doc.close()
    return output_path


def write_document_to_disk(
    html_content: str,
    output_path: Path,
    extension: str,
    font_info: Dict[str, Any],
    pages_data: Optional[List[Dict[str, Any]]] = None,
    original_file_path: Optional[Path] = None,
    work_dir: Optional[Path] = None,
) -> Path:
    """
    Exports the edited document to output_path matching the chosen extension.
    If pages_data is present, exports in layout-preserved mode so tables, lines,
    borders, graphics, and separate pages are 100% maintained.
    """
    html_content = sanitize_document_html(html_content) if html_content else ""
    ext = extension.lstrip(".").lower()
    kind = EXTENSION_KINDS.get(ext, "text")

    font_name = font_info.get("name", "EB Garamond")
    font_path = font_info.get("font_path")
    font_size_pt = float(font_info.get("size_pt", 12.0))

    # If layout-preserved mode with pages_data
    if pages_data and len(pages_data) > 0:
        if kind == "searchable_pdf" or ext == "searchable_pdf":
            return _write_searchable_pdf(
                pages_data=pages_data,
                work_dir=work_dir,
                original_file_path=original_file_path,
                output_path=output_path,
                font_name=font_name,
                font_path=font_path,
                font_size_pt=font_size_pt,
            )
        elif kind == "pdf":
            return _write_pdf_layout_preserved(
                pages_data=pages_data,
                original_file_path=original_file_path,
                output_path=output_path,
                font_name=font_name,
                font_path=font_path,
                font_size_pt=font_size_pt,
                work_dir=work_dir,
            )
        elif kind == "image":
            return _write_image_layout_preserved(
                pages_data=pages_data,
                work_dir=work_dir,
                output_path=output_path,
                image_ext=ext,
                font_name=font_name,
                font_path=font_path,
                font_size_pt=font_size_pt,
            )

    if kind in ("pdf", "searchable_pdf") or ext == "searchable_pdf":
        return _write_pdf(html_content, output_path, font_name, font_path, font_size_pt, original_file_path=original_file_path)
    elif kind == "image":
        return _write_image(html_content, output_path, ext, font_name, font_path, font_size_pt)
    elif kind == "docx":
        return _write_docx(html_content, output_path, font_name, font_size_pt, original_file_path=original_file_path)
    elif kind == "markdown":
        return _write_markdown(html_content, output_path)
    elif kind == "html":
        return _write_html(html_content, output_path, font_name, font_size_pt)
    else:  # Text default
        return _write_text(html_content, output_path)



def _prepare_clean_html_for_story(raw_html: str, font_name: str, font_size_pt: float) -> str:
    """
    Prepares HTML body for PyMuPDF Story rendering with full table and line CSS.
    """
    safe_font = sanitize_font_name(font_name)
    sanitized = sanitize_document_html(raw_html)
    cleaned = re.sub(r'<span[^>]*class="[^"]*docedit-fuzzy[^"]*"[^>]*>(.*?)</span>', r'\1', sanitized)
    cleaned = re.sub(r'data-[a-zA-Z0-9_\-]+="[^"]*"', '', cleaned)
    body = f"""
    <div style="font-family: '{safe_font}', serif; font-size: {font_size_pt}pt; line-height: 1.28; color: #111111;">
        {cleaned}
    </div>
    """
    return body


def _write_pdf(
    html_content: str,
    output_path: Path,
    font_name: str,
    font_path: Optional[str],
    font_size_pt: float,
    original_file_path: Optional[Path] = None,
) -> Path:
    """Generates PDF using PyMuPDF Story with full table and vector line preservation."""
    import pymupdf

    safe_font = sanitize_font_name(font_name)
    story_html = _prepare_clean_html_for_story(html_content, safe_font, font_size_pt)

    # Base styling for tables, lines, headings, and paragraphs
    base_css = f"""
    body, div {{
        font-family: '{safe_font}', serif;
        font-size: {font_size_pt}pt;
        line-height: 1.28;
        color: #111111;
    }}
    table, .docedit-table {{
        border-collapse: collapse;
        width: 100%;
        margin: 8pt 0;
        border: 1pt solid #555555;
    }}
    th, td {{
        border: 1pt solid #666666;
        padding: 4pt 7pt;
        text-align: left;
        font-size: {max(8.0, font_size_pt - 1.5)}pt;
    }}
    th {{
        font-weight: bold;
        background-color: #ededed;
    }}
    hr, .docedit-line {{
        border: none;
        border-top: 1pt solid #666666;
        margin: 8pt 0;
    }}
    h1 {{ font-size: {font_size_pt + 5}pt; font-weight: bold; margin: 4pt 0 3pt 0; text-align: center; }}
    h2 {{ font-size: {font_size_pt + 3}pt; font-weight: bold; margin: 7pt 0 3pt 0; text-align: center; }}
    h3 {{ font-size: {font_size_pt + 1.5}pt; font-weight: bold; margin: 5pt 0 2pt 0; }}
    p {{ margin: 0 0 2.5pt 0; line-height: 1.25; }}
    ul {{ margin: 2pt 0 4pt 0; padding-left: 18pt; }}
    li {{ margin: 1.5pt 0; line-height: 1.25; }}
    .docedit-docx-icon {{ height: 1.1em; max-height: 16px; vertical-align: -0.15em; margin: 0 3px; display: inline-block; }}
    a {{ color: inherit; text-decoration: underline; }}
    """

    custom_css = base_css
    archive = None
    if font_path and Path(font_path).is_file():
        fp = Path(font_path)
        archive = pymupdf.Archive(str(fp.parent))
        custom_css += f"""
        @font-face {{
            font-family: '{safe_font}';
            src: url('{fp.name}');
        }}
        """

    story = pymupdf.Story(story_html, user_css=custom_css, archive=archive)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    page_rect = pymupdf.paper_rect("a4")

    # Adapt margin to match source document if DOCX (e.g. 42.5 pt margins)
    margin = 42.5  # ~0.59 in compact margins
    if original_file_path and Path(original_file_path).is_file() and str(original_file_path).lower().endswith(".docx"):
        try:
            import docx
            orig_doc = docx.Document(str(original_file_path))
            if orig_doc.sections:
                sec = orig_doc.sections[0]
                margin = min(sec.top_margin.pt, sec.left_margin.pt, 45.0)
        except Exception:
            pass

    content_rect = page_rect + (margin, margin, -margin, -margin)

    more = 1
    while more:
        dev = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(dev)
        writer.end_page()
    writer.close()

    output_path.write_bytes(buf.getvalue())
    return output_path


def _write_image(
    html_content: str,
    output_path: Path,
    image_ext: str,
    font_name: str,
    font_path: Optional[str],
    font_size_pt: float
) -> Path:
    """Renders the edited document with tables & lines to an image using PyMuPDF pixmap."""
    import pymupdf
    from PIL import Image

    pdf_tmp = output_path.with_suffix(".tmp.pdf")
    _write_pdf(html_content, pdf_tmp, font_name, font_path, font_size_pt)

    doc = pymupdf.open(str(pdf_tmp))
    if len(doc) == 0:
        pdf_tmp.unlink(missing_ok=True)
        return output_path

    fmt_map = {
        "jpg": "JPEG",
        "jpeg": "JPEG",
        "png": "PNG",
        "tif": "TIFF",
        "tiff": "TIFF",
        "bmp": "BMP",
        "webp": "WEBP",
    }
    save_fmt = fmt_map.get(image_ext.lower(), "PNG")

    pix = doc[0].get_pixmap(dpi=conf.OCR_DPI)
    Image.MAX_IMAGE_PIXELS = getattr(conf, "MAX_IMAGE_PIXELS", 50_000_000)
    img = Image.open(io.BytesIO(pix.tobytes("png")))

    if save_fmt == "JPEG" and img.mode in ("RGBA", "P"):
        img = img.convert("RGB")

    if save_fmt == "TIFF" and len(doc) > 1:
        other_frames = []
        for p_idx in range(1, len(doc)):
            p_pix = doc[p_idx].get_pixmap(dpi=conf.OCR_DPI)
            p_img = Image.open(io.BytesIO(p_pix.tobytes("png")))
            other_frames.append(p_img)
        img.save(str(output_path), format=save_fmt, save_all=True, append_images=other_frames, dpi=(conf.OCR_DPI, conf.OCR_DPI))
    else:
        img.save(str(output_path), format=save_fmt, dpi=(conf.OCR_DPI, conf.OCR_DPI))

    doc.close()
    pdf_tmp.unlink(missing_ok=True)
    return output_path


def _write_docx(
    html_content: str,
    output_path: Path,
    font_name: str,
    font_size_pt: float,
    original_file_path: Optional[Path] = None,
) -> Path:
    """
    Generates Word DOCX preserving full table structures, header shading,
    borders, divider lines, headings, alignments, matching typography, inline icons,
    and hyperlinks.
    """
    import base64
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Inches, Pt

    blocks = parse_html_blocks(html_content)

    doc = docx.Document()

    # Set normal style typography
    style = doc.styles["Normal"]
    font = style.font
    font.name = font_name
    font.size = Pt(font_size_pt)
    try:
        rPr_norm = style.element.get_or_add_rPr()
        rFonts_norm = parse_xml(r'<w:rFonts {} w:ascii="{}" w:hAnsi="{}" w:cs="{}"/>'.format(nsdecls('w'), font_name, font_name, font_name))
        rPr_norm.append(rFonts_norm)
    except Exception:
        pass

    # Section margins: match original DOCX if available, else default to compact 0.55in/0.6in margins
    if original_file_path and Path(original_file_path).is_file() and str(original_file_path).lower().endswith(".docx"):
        try:
            orig_doc = docx.Document(str(original_file_path))
            if orig_doc.sections:
                orig_sec = orig_doc.sections[0]
                for section in doc.sections:
                    section.top_margin = orig_sec.top_margin
                    section.bottom_margin = orig_sec.bottom_margin
                    section.left_margin = orig_sec.left_margin
                    section.right_margin = orig_sec.right_margin
                    section.page_width = orig_sec.page_width
                    section.page_height = orig_sec.page_height
        except Exception as e:
            logger.debug("Failed copying original docx section margins: %s", e)
    else:
        for section in doc.sections:
            section.top_margin = Inches(0.55)
            section.bottom_margin = Inches(0.55)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

    align_map = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }

    def add_formatted_runs(p, raw_text, base_size=font_size_pt, default_run_font=font_name):
        from docx.shared import RGBColor
        tokens = re.split(r'(<[^>]+>)', raw_text)
        bold = False
        italic = False
        underline = False
        font_stack: List[Optional[str]] = []
        size_stack: List[Optional[float]] = []
        color_stack: List[Optional[RGBColor]] = []
        active_url: Optional[str] = None

        for tok in tokens:
            if not tok:
                continue
            lower_tok = tok.lower()
            if lower_tok in ('<b>', '<strong>'):
                bold = True
            elif lower_tok in ('</b>', '</strong>'):
                bold = False
            elif lower_tok in ('<i>', '<em>'):
                italic = True
            elif lower_tok in ('</i>', '</em>'):
                italic = False
            elif lower_tok == '<u>':
                underline = True
            elif lower_tok == '</u>':
                underline = False
            elif lower_tok in ('<br>', '<br/>', '<br />'):
                p.add_run('\n')
            elif lower_tok.startswith("<img") and lower_tok.endswith(">"):
                src_m = re.search(r'src=["\']([^"\']+)["\']', tok, re.IGNORECASE)
                if src_m:
                    src_val = src_m.group(1)
                    try:
                        raw_img = None
                        if src_val.startswith("data:image/"):
                            raw_img = base64.b64decode(src_val.split(",", 1)[1])
                        elif Path(src_val).is_file():
                            raw_img = Path(src_val).read_bytes()
                        if raw_img:
                            r_img = p.add_run()
                            # Determine dimensions (default icon ~11pt)
                            w_match = re.search(r'width:\s*([\d\.]+)px', tok, re.I)
                            h_match = re.search(r'height:\s*([\d\.]+)px', tok, re.I)
                            w_pt = float(w_match.group(1)) * 0.75 if w_match else 11.0
                            h_pt = float(h_match.group(1)) * 0.75 if h_match else 11.0
                            r_img.add_picture(io.BytesIO(raw_img), width=Pt(w_pt), height=Pt(h_pt))
                    except Exception as e:
                        logger.debug("Failed adding picture to docx run: %s", e)
            elif lower_tok.startswith("<a ") and lower_tok.endswith(">"):
                href_m = re.search(r'href=["\']([^"\']+)["\']', tok, re.IGNORECASE)
                if href_m:
                    active_url = href_m.group(1)
            elif lower_tok == "</a>":
                active_url = None
            elif lower_tok.startswith('<span') and lower_tok.endswith('>'):
                f_match = re.search(r'font-family:\s*[\'"]?([^\'";]+)[\'"]?', tok, re.IGNORECASE)
                s_match = re.search(r'font-size:\s*([\d\.]+)pt', tok, re.IGNORECASE)
                c_match = re.search(r'color:\s*#([0-9a-fA-F]{6})', tok, re.IGNORECASE)
                cur_f = f_match.group(1).split(",")[0].strip().strip("'\"") if f_match else None
                cur_s = float(s_match.group(1)) if s_match else None
                cur_c = None
                if c_match:
                    h_val = c_match.group(1)
                    cur_c = RGBColor(int(h_val[0:2], 16), int(h_val[2:4], 16), int(h_val[4:6], 16))
                font_stack.append(cur_f)
                size_stack.append(cur_s)
                color_stack.append(cur_c)
            elif lower_tok == '</span>':
                if font_stack:
                    font_stack.pop()
                if size_stack:
                    size_stack.pop()
                if color_stack:
                    color_stack.pop()
            elif tok.startswith('<') and tok.endswith('>'):
                pass  # Ignore other structural tags in run text
            else:
                clean = html.unescape(tok)
                if clean:
                    active_font = next((f for f in reversed(font_stack) if f), default_run_font)
                    active_size = next((s for s in reversed(size_stack) if s is not None), base_size)
                    active_color = next((c for c in reversed(color_stack) if c is not None), None)

                    if active_url:
                        try:
                            part = p.part
                            r_id = part.relate_to(active_url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                            safe_txt = html.escape(clean)
                            c_val = "0563C1"
                            if active_color:
                                c_val = f"{active_color[0]:02x}{active_color[1]:02x}{active_color[2]:02x}"
                            hl_xml = (
                                f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                                f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}">'
                                f'<w:r><w:rPr><w:rFonts w:ascii="{active_font}" w:hAnsi="{active_font}"/>'
                                f'<w:sz w:val="{int(active_size * 2)}"/><w:u w:val="single"/><w:color w:val="{c_val}"/></w:rPr>'
                                f'<w:t>{safe_txt}</w:t></w:r></w:hyperlink>'
                            )
                            p._p.append(parse_xml(hl_xml))
                        except Exception:
                            r = p.add_run(clean)
                            r.font.name = active_font
                            r.font.size = Pt(active_size)
                            r.underline = True
                    else:
                        r = p.add_run(clean)
                        r.font.name = active_font
                        try:
                            rPr = r._r.get_or_add_rPr()
                            rFonts = parse_xml(r'<w:rFonts {} w:ascii="{}" w:hAnsi="{}" w:cs="{}"/>'.format(nsdecls('w'), active_font, active_font, active_font))
                            rPr.append(rFonts)
                        except Exception:
                            pass
                        r.font.size = Pt(active_size)
                        if active_color:
                            r.font.color.rgb = active_color
                        r.bold = bold
                        r.italic = italic
                        r.underline = underline

    for block in blocks:
        b_type = block["type"]

        if b_type == "hr":
            # Native divider line in Word via border
            p_hr = doc.add_paragraph()
            p_hr.paragraph_format.space_before = Pt(6)
            p_hr.paragraph_format.space_after = Pt(6)
            pPr = p_hr._p.get_or_add_pPr()
            pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="8" w:space="1" w:color="777777"/></w:pBdr>')
            pPr.append(pBdr)

        elif b_type in ("h1", "h2", "h3"):
            level = int(b_type[1])
            p_h = doc.add_heading(level=level)
            p_h.paragraph_format.space_before = Pt(6)
            p_h.paragraph_format.space_after = Pt(2)
            align = block.get("align")
            if align in align_map:
                p_h.alignment = align_map[align]
            h_size = block.get("font_size_pt") or (font_size_pt + (6 - level * 2))
            b_font = block.get("font_name") or font_name
            add_formatted_runs(p_h, block["html"], base_size=h_size, default_run_font=b_font)

        elif b_type == "table":
            rows = block["rows"]
            if not rows:
                continue
            col_count = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=col_count)
            table.style = "Table Grid"

            for r_idx, row in enumerate(rows):
                for c_idx, cell_data in enumerate(row):
                    if c_idx >= col_count:
                        break
                    cell = table.cell(r_idx, c_idx)
                    p_cell = cell.paragraphs[0]
                    p_cell.paragraph_format.space_after = Pt(2)
                    p_cell.paragraph_format.space_before = Pt(2)

                    is_hdr = cell_data.get("is_header", False)
                    if is_hdr:
                        shd = parse_xml(r'<w:shd {} w:fill="EAEAEA"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shd)

                    cell_html = cell_data.get("html", "")
                    if is_hdr:
                        cell_html = f"<strong>{cell_html}</strong>"
                    add_formatted_runs(p_cell, cell_html, base_size=max(9.0, font_size_pt - 1.5))

            # Add breathing room after table
            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_before = Pt(3)
            p_spacer.paragraph_format.space_after = Pt(3)

        elif b_type == "li":
            p_li = doc.add_paragraph(style="List Bullet")
            p_li.paragraph_format.space_before = Pt(0)
            p_li.paragraph_format.space_after = Pt(1.5)
            p_li.paragraph_format.line_spacing = 1.1
            b_size = block.get("font_size_pt") or font_size_pt
            b_font = block.get("font_name") or font_name
            add_formatted_runs(p_li, block["html"], base_size=b_size, default_run_font=b_font)

        else:  # Paragraph 'p' or blockquote
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            align = block.get("align")
            if align in align_map:
                p.alignment = align_map[align]
            b_size = block.get("font_size_pt") or font_size_pt
            b_font = block.get("font_name") or font_name
            add_formatted_runs(p, block["html"], base_size=b_size, default_run_font=b_font)

    doc.save(str(output_path))
    return output_path


def _write_markdown(html_content: str, output_path: Path) -> Path:
    """Converts HTML blocks to clean markdown preserving tables and divider lines."""
    blocks = parse_html_blocks(html_content)
    md_lines = []

    for block in blocks:
        b_type = block["type"]
        if b_type == "hr":
            md_lines.append("\n---\n")
        elif b_type == "h1":
            text = html.unescape(re.sub(r'<[^>]+>', '', block["html"]))
            md_lines.append(f"# {text}\n")
        elif b_type == "h2":
            text = html.unescape(re.sub(r'<[^>]+>', '', block["html"]))
            md_lines.append(f"## {text}\n")
        elif b_type == "h3":
            text = html.unescape(re.sub(r'<[^>]+>', '', block["html"]))
            md_lines.append(f"### {text}\n")
        elif b_type == "li":
            text = html.unescape(re.sub(r'<[^>]+>', '', block["html"]))
            md_lines.append(f"- {text}")
        elif b_type == "table":
            rows = block["rows"]
            if not rows:
                continue
            col_count = max(len(r) for r in rows)
            table_md = []
            # Header row
            header_cells = [html.unescape(re.sub(r'<[^>]+>', '', cell.get("html", ""))) for cell in rows[0]]
            while len(header_cells) < col_count:
                header_cells.append("")
            table_md.append("| " + " | ".join(header_cells) + " |")
            table_md.append("| " + " | ".join(["---"] * col_count) + " |")

            for row in rows[1:]:
                cells = [html.unescape(re.sub(r'<[^>]+>', '', cell.get("html", ""))) for cell in row]
                while len(cells) < col_count:
                    cells.append("")
                table_md.append("| " + " | ".join(cells) + " |")

            md_lines.append("\n" + "\n".join(table_md) + "\n")
        else:  # Paragraph 'p'
            text = block["html"]
            text = re.sub(r'<strong>(.*?)</strong>', r'**\1**', text, flags=re.IGNORECASE)
            text = re.sub(r'<em>(.*?)</em>', r'*\1*', text, flags=re.IGNORECASE)
            text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
            clean = html.unescape(re.sub(r'<[^>]+>', '', text)).strip()
            if clean:
                md_lines.append(f"{clean}\n")

    output_path.write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
    return output_path


def _write_html(html_content: str, output_path: Path, font_name: str, font_size_pt: float) -> Path:
    """Wraps HTML in clean HTML5 document with embedded styling for tables, lines & typography."""
    safe_font = sanitize_font_name(font_name)
    cleaned = sanitize_document_html(html_content)
    cleaned = re.sub(r'<span[^>]*class="[^"]*docedit-fuzzy[^"]*"[^>]*>(.*?)</span>', r'\1', cleaned)
    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Edited Document</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family={safe_font.replace(' ', '+')}:wght@400;700&display=swap">
<style>
body {{
    font-family: '{safe_font}', serif;
    font-size: {font_size_pt}pt;
    line-height: 1.6;
    color: #212529;
    max-width: 860px;
    margin: 2.5rem auto;
    padding: 0 1.5rem;
}}
table, .docedit-table {{
    border-collapse: collapse;
    width: 100%;
    margin: 1.5rem 0;
    border: 1px solid #777777;
}}
th, td {{
    border: 1px solid #777777;
    padding: 8px 12px;
    text-align: left;
}}
th {{
    font-weight: bold;
    background-color: #f2f2f2;
}}
hr, .docedit-line {{
    border: none;
    border-top: 1px solid #777777;
    margin: 1.5rem 0;
}}
h1, h2, h3 {{ margin-top: 1.5rem; margin-bottom: 0.6rem; }}
p {{ margin-bottom: 1.2rem; text-align: justify; }}
</style>
</head>
<body>
{cleaned}
</body>
</html>
"""
    output_path.write_text(full_html, encoding="utf-8")
    return output_path


def _write_text(html_content: str, output_path: Path) -> Path:
    """Formats HTML to clean plain text with structured ASCII tables and divider rules."""
    blocks = parse_html_blocks(html_content)
    text_lines = []

    for block in blocks:
        b_type = block["type"]
        if b_type == "hr":
            text_lines.append("-" * 60)
        elif b_type in ("h1", "h2", "h3"):
            txt = html.unescape(re.sub(r'<[^>]+>', '', block["html"])).strip()
            text_lines.append(txt)
            text_lines.append("=" * min(len(txt), 60) if b_type == "h1" else "-" * min(len(txt), 60))
        elif b_type == "li":
            txt = html.unescape(re.sub(r'<[^>]+>', '', block["html"])).strip()
            text_lines.append(f"  * {txt}")
        elif b_type == "table":
            rows = block["rows"]
            if not rows:
                continue
            # Format as clean ASCII table
            table_matrix = []
            col_count = max(len(r) for r in rows)
            for row in rows:
                row_vals = [html.unescape(re.sub(r'<[^>]+>', '', c.get("html", ""))).strip() for c in row]
                while len(row_vals) < col_count:
                    row_vals.append("")
                table_matrix.append(row_vals)

            col_widths = [max(len(table_matrix[r][c]) for r in range(len(table_matrix))) for c in range(col_count)]
            col_widths = [max(w, 4) for w in col_widths]

            # Render rows
            for r_idx, row in enumerate(table_matrix):
                cells_str = [row[c].ljust(col_widths[c]) for c in range(col_count)]
                text_lines.append(" | ".join(cells_str))
                if r_idx == 0:
                    text_lines.append("-+-".join("-" * col_widths[c] for c in range(col_count)))
            text_lines.append("")
        else:  # Paragraph
            raw = re.sub(r'<br\s*/?>', '\n', block["html"], flags=re.IGNORECASE)
            txt = html.unescape(re.sub(r'<[^>]+>', '', raw)).strip()
            if txt:
                text_lines.append(txt)
                text_lines.append("")

    output_path.write_text("\n".join(text_lines).strip() + "\n", encoding="utf-8")
    return output_path
